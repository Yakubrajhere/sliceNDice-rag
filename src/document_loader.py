"""Load PDF, TXT, DOCX (and MD) files from a directory into plain-text documents."""
from pathlib import Path
from typing import Dict, List

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


def _load_pdf(path: Path) -> List[Dict]:
    """One document per PDF page so page numbers survive into citations."""
    docs = []
    reader = PdfReader(str(path))
    for page_num, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append({"text": text, "metadata": {"source": path.name, "page": page_num}})
    return docs


def _load_txt(path: Path) -> List[Dict]:
    text = path.read_text(encoding="utf-8", errors="ignore").strip()
    return [{"text": text, "metadata": {"source": path.name, "page": 1}}] if text else []


def _load_docx(path: Path) -> List[Dict]:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return [{"text": text, "metadata": {"source": path.name, "page": 1}}] if text else []


_LOADERS = {
    ".pdf": _load_pdf,
    ".txt": _load_txt,
    ".md": _load_txt,
    ".docx": _load_docx,
}


def load_file(path: Path) -> List[Dict]:
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        return []
    return loader(path)


def load_documents(data_dir: str) -> List[Dict]:
    """Load every supported file under data_dir (recursive)."""
    root = Path(data_dir)
    if not root.exists():
        raise FileNotFoundError(f"Data directory not found: {root.resolve()}")

    documents: List[Dict] = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                loaded = load_file(path)
                print(f"  loaded {path.name}: {len(loaded)} section(s)")
                documents.extend(loaded)
            except Exception as exc:  # keep going if one file is bad
                print(f"  [skip] {path.name}: {exc}")
    return documents
